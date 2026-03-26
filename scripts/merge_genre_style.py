#!/usr/bin/env python3
"""
merge_genre_style.py
====================
장르 + 스타일 계열을 하나의 '장르' 카테고리로 통합하고,
하위 그룹(Photography, Film & Video, Animation & Cartoon, Anime, 
Illustration & Art, Game, Style)으로 자동 분류합니다.

- Token_DB: 스타일 계열 → 장르로 category 변경, group 재배치
- Group_DB: 스타일 계열 행 삭제, 장르 하위 그룹 행 추가
"""

import openpyxl
import shutil
from pathlib import Path

BASE_DIR = Path(__file__).parent.parent
EXCEL_PATH = BASE_DIR / "data" / "prompt_master_v2.xlsx"

# ─────────────────────────────────────────────
# Group classification rules
# ─────────────────────────────────────────────
GENRE_CAT_KEY = "cat_7e0c0a64"

GROUP_RULES = {
    "Photography": [
        "Portrait Photography", "Landscape Photography", "Street Photography",
        "Night Photography", "Aerial Photography", "Architectural Photography",
        "Documentary Photography", "Fashion Editorials", "Food Photography",
        "Macro Photography", "Product Photography", "Sports/Action Photography",
        "Underwater Photography", "Milky Way/Star Photography", "Birds/Wildlife",
        "Concerts/Performances", "fashion editorial photography", "portrait photography",
    ],
    "Film & Video": [
        "cinematic film still", "commercial film frame", "animation film frame",
        "Basic Videography", "indie movie animation",
    ],
    "Animation & Cartoon": [
        "2D animation still", "3D cartoon style", "cartoon style", "cartoonish 3D",
        "Disney Renaissance style", "DreamWorks", "Illumination studio",
        "Pixar style", "Pixar/Blender style", "Pixar animation concept art style",
        "Hyper realistic Pixar style 3D character",
        "A line sketch of a disney-style drawing",
    ],
    "Anime": [
        "Japanese anime", "anime-style", "Makoto Shinkai style",
        "Japanese anime illustration style",
    ],
    "Illustration & Art": [
        "concept art", "digital illustration", "digital painting",
        "illustration", "minimalist illustration", "oil painting", "watercolo",
        "fantasy art style", "A stunning fantasy portrait",
        "DnD art style", "Dungeons and Dragons art style", "dc comics",
    ],
    "Game": [
        "3D game asset", "Game asset design", "game character design", "game environment design",
        "casual game", "casual mobile game art style", "Gardenescapes",
        "Overwatch", "fortnite style",
        "bright western casual mobile game style", "colorful social mobile game style",
    ],
    "Style": [
        "photorealistic", "semi-realistic character design", "stylized 3D render",
    ],
}

# Build reverse lookup: value -> group_name
VALUE_TO_GROUP = {}
for group_name, values in GROUP_RULES.items():
    for v in values:
        VALUE_TO_GROUP[v] = group_name


def clean(val):
    if val is None:
        return ""
    return str(val).strip()


def classify_token(token_value):
    """Classify a token_value into a group name."""
    # Direct match
    if token_value in VALUE_TO_GROUP:
        return VALUE_TO_GROUP[token_value]
    # Case-insensitive match
    for v, g in VALUE_TO_GROUP.items():
        if v.lower() == token_value.lower():
            return g
    # Partial match heuristics
    tv_lower = token_value.lower()
    if "photo" in tv_lower or "wildlife" in tv_lower:
        return "Photography"
    if "film" in tv_lower or "video" in tv_lower or "cinematic" in tv_lower:
        return "Film & Video"
    if "anime" in tv_lower:
        return "Anime"
    if "cartoon" in tv_lower or "pixar" in tv_lower or "disney" in tv_lower or "dreamworks" in tv_lower or "illumination" in tv_lower:
        return "Animation & Cartoon"
    if "game" in tv_lower or "fortnite" in tv_lower or "overwatch" in tv_lower:
        return "Game"
    if "illustration" in tv_lower or "painting" in tv_lower or "art" in tv_lower or "comics" in tv_lower:
        return "Illustration & Art"
    # Fallback
    return "Style"


def main():
    print("=" * 60)
    print("Merge 장르 + 스타일 계열 → 장르 (with subgroups)")
    print("=" * 60)

    # Backup
    backup = EXCEL_PATH.with_suffix(EXCEL_PATH.suffix + ".pre-merge-backup")
    shutil.copy2(EXCEL_PATH, backup)
    print(f"Backup: {backup.name}")

    wb = openpyxl.load_workbook(EXCEL_PATH)

    # ── 1. Update Token_DB ──
    print("\n[1/3] Updating Token_DB...")
    ws_db = wb["Token_DB"]
    updated = 0
    reclassified = 0

    for row in range(2, ws_db.max_row + 1):
        cat = clean(ws_db.cell(row, 3).value)      # category
        tval = clean(ws_db.cell(row, 10).value)     # token_value
        short = clean(ws_db.cell(row, 7).value)     # token_label_short

        if cat not in ("장르", "스타일 계열"):
            continue

        # Use short label or token_value for classification
        match_val = short or tval
        new_group = classify_token(match_val)

        # Update category to 장르 (if 스타일 계열)
        if cat == "스타일 계열":
            ws_db.cell(row, 3, "장르")              # category
            ws_db.cell(row, 4, GENRE_CAT_KEY)       # category_key
            # Update token_id prefix if needed
            old_id = clean(ws_db.cell(row, 1).value)
            if old_id.startswith("cat_518ce1f8"):
                new_id = old_id.replace("cat_518ce1f8", GENRE_CAT_KEY)
                ws_db.cell(row, 1, new_id)
            reclassified += 1

        # Update group
        ws_db.cell(row, 5, new_group)               # group
        updated += 1

    print(f"   {updated} rows updated ({reclassified} moved from 스타일 계열)")

    # ── 2. Update Group_DB ──
    print("\n[2/3] Updating Group_DB...")
    ws_grp = wb["Group_DB"]

    # Find and remove existing 장르 and 스타일 계열 rows
    rows_to_delete = []
    for row in range(2, ws_grp.max_row + 1):
        cat = clean(ws_grp.cell(row, 1).value)
        if cat in ("장르", "스타일 계열"):
            rows_to_delete.append(row)

    # Delete in reverse order
    for row in sorted(rows_to_delete, reverse=True):
        ws_grp.delete_rows(row)
        print(f"   Deleted Group_DB row {row}")

    # Add new group rows for 장르
    GROUP_ORDER = [
        ("Photography", 1),
        ("Film & Video", 2),
        ("Animation & Cartoon", 3),
        ("Anime", 4),
        ("Illustration & Art", 5),
        ("Game", 6),
        ("Style", 7),
    ]

    next_row = ws_grp.max_row + 1
    for group_name, sort_order in GROUP_ORDER:
        ws_grp.cell(next_row, 1, "장르")                # category
        ws_grp.cell(next_row, 2, group_name)            # group
        ws_grp.cell(next_row, 3, group_name)            # group_display_name
        ws_grp.cell(next_row, 4, "")                    # notes
        ws_grp.cell(next_row, 5, sort_order)            # sort_order
        print(f"   Added: 장르 / {group_name} (sort: {sort_order})")
        next_row += 1

    # ── 3. Update Token_Input ──
    print("\n[3/3] Updating Token_Input...")
    ws_input = wb["Token_Input"]
    input_updated = 0

    for row in range(2, ws_input.max_row + 1):
        cat = clean(ws_input.cell(row, 3).value)
        tval = clean(ws_input.cell(row, 7).value)
        name = clean(ws_input.cell(row, 5).value)

        if cat not in ("장르", "스타일 계열"):
            continue

        match_val = name or tval
        new_group = classify_token(match_val)

        if cat == "스타일 계열":
            ws_input.cell(row, 3, "장르")

        ws_input.cell(row, 4, new_group)
        input_updated += 1

    print(f"   {input_updated} Token_Input rows updated")

    # Save
    wb.save(EXCEL_PATH)
    print(f"\n{'=' * 60}")
    print("Excel saved!")
    print(f"   Token_DB: {updated} rows (장르 subgroups)")
    print(f"   Group_DB: {len(GROUP_ORDER)} new groups for 장르")
    print(f"   Token_Input: {input_updated} rows")
    print(f"\nNext: run 'python scripts/convert_excel_v2.py'")


if __name__ == "__main__":
    main()
