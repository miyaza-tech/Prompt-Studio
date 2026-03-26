#!/usr/bin/env python3
"""
Add option descriptions sheet to Excel
"""

import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

BASE_DIR = Path(__file__).parent.parent
EXCEL_PATH = BASE_DIR / "data" / "prompt_master.xlsx"

# 샘플 설명 데이터 (카테고리별 주요 옵션)
DESCRIPTIONS = {
    # 샷 크기
    "Extreme Close-Up (익스트림 클로즈업)": "눈, 입술 등 얼굴 일부만 화면 가득 채움",
    "Close-Up (클로즈업)": "얼굴 전체를 프레임에 담는 샷",
    "Medium Close-Up (미디엄 클로즈업)": "어깨 위로 얼굴까지 담는 샷",
    "Medium Shot (미디엄 샷)": "허리 위로 상체를 담는 샷",
    "Medium Long Shot (미디엄 롱 샷)": "무릎 위로 신체를 담는 샷",
    "Long Shot (롱 샷)": "전신을 담되 주변 공간도 보이는 샷",
    "Extreme Long Shot (익스트림 롱 샷)": "인물이 작게 보이고 배경이 주가 됨",
    "Wide Shot (와이드 샷)": "전체 장면과 환경을 넓게 담는 샷",
    "Extreme Wide Shot (익스트림 와이드 샷)": "매우 넓은 풍경이나 도시 전경",
    "Full Shot (풀 샷)": "인물 전신이 프레임에 꽉 차게",
    "Two-Shot (투 샷)": "두 인물을 함께 담는 구도",
    
    # 샷 앵글
    "Eye Level Shot (눈높이 샷)": "피사체와 같은 높이에서 촬영",
    "higher angle (하이 앵글)": "위에서 아래로 내려다보는 각도",
    "Low Angle (로우 앵글)": "아래에서 위로 올려다보는 각도",
    "Bird's Eye View (조감도)": "새가 하늘에서 내려다보는 시점",
    "Worm's Eye View (벌레 시점)": "바닥에서 극단적으로 올려다보는 시점",
    "Dutch Angle (더치 앵글)": "카메라를 기울여 불안정한 느낌",
    "Over-the-Shoulder (오버 더 숄더)": "어깨 너머로 상대방을 보는 구도",
    "POV (1인칭 시점)": "캐릭터의 눈으로 보는 시점",
    
    # 조명
    "Soft Light (소프트 라이트)": "부드럽고 그림자가 약한 조명",
    "Hard Light (하드 라이트)": "강하고 선명한 그림자를 만드는 조명",
    "Natural Light (자연광)": "창문이나 야외의 자연스러운 빛",
    "Studio Light (스튜디오 조명)": "인공적으로 세팅된 스튜디오 조명",
    "Rembrandt Lighting (렘브란트 조명)": "얼굴 한쪽에 삼각형 빛이 생기는 기법",
    "Butterfly Lighting (버터플라이 조명)": "코 아래 나비 모양 그림자 생성",
    "Loop Lighting (루프 조명)": "코 옆에 작은 그림자가 생기는 조명",
    "Split Lighting (스플릿)": "얼굴 절반만 밝히는 극적인 조명",
    "Backlighting (역광)": "피사체 뒤에서 오는 빛",
    "Rim Light (림 라이트)": "피사체 윤곽선만 빛나게 하는 조명",
    "Golden Hour (골든 아워)": "일출/일몰 시간의 황금빛 자연광",
    "Cinematic Lighting (영화적 조명)": "영화처럼 드라마틱한 조명 설정",
    "Volumetric Lighting (볼류메트릭)": "빛줄기가 보이는 대기 효과",
    
    # 색채
    "Soft color (부드러운 색상)": "채도가 낮고 파스텔톤의 색상",
    "Vibrant Colors (선명한 색상)": "채도가 높고 생동감 있는 색상",
    "Muted Colors (차분한 색상)": "채도를 낮춘 차분한 톤",
    "Monochromatic (단색)": "한 가지 색의 다양한 명도",
    "Warm Tones (따뜻한 톤)": "노랑, 주황, 빨강 계열의 따뜻한 색",
    "Cool Tones (차가운 톤)": "파랑, 녹색 계열의 차가운 색",
    "Pastel Colors (파스텔 색상)": "연하고 부드러운 파스텔톤",
    "Neon Colors (네온 색상)": "밝고 형광빛 나는 네온 색상",
    "Black and White (흑백)": "컬러 없이 흑백으로만 표현",
    
    # 분위기
    "Moody (무디)": "어둡고 감성적인 분위기",
    "Dreamy (몽환적)": "꿈처럼 흐릿하고 환상적인 느낌",
    "Ethereal (에테리얼)": "천상의, 숭고한 분위기",
    "Gritty (그리티)": "거칠고 현실적인 질감",
    "Minimalist (미니멀리스트)": "단순하고 여백이 많은 구성",
    "Dramatic (드라마틱)": "극적이고 강렬한 분위기",
    "Nostalgic (노스탤직)": "향수를 불러일으키는 복고풍",
    "Mysterious (신비로운)": "신비롭고 알 수 없는 분위기",
    "Peaceful (평화로운)": "고요하고 평화로운 느낌",
    "Epic (장엄한)": "웅장하고 스케일이 큰 느낌",
    
    # 렌즈 기법
    "Bokeh (보케)": "배경을 흐리게 처리하는 효과",
    "Deep Focus (딥 포커스)": "전경부터 배경까지 모두 선명",
    "Shallow DOF (얀 심도 / 보케)": "피사체만 선명하고 배경 흐림",
    "Rack Focus (랙 포커스)": "초점을 이동시키는 기법",
    "Soft Focus (소프트 포커스)": "전체적으로 부드럽게 흐린 효과",
    "Lens Flare (렌즈 플레어)": "렌즈에 빛이 반사되는 효과",
    "Long Exposure (장노출)": "빛의 궤적이 보이는 장노출",
    "Double Exposure (이중 노출)": "두 이미지를 겹쳐서 합성",
    "Vignette (비네트)": "사진 가장자리를 어둡게 처리",
    
    # 구도
    "Rule of Thirds (3분할 법칙)": "화면을 9등분하여 교차점에 배치",
    "Center Framing (중앙 배치)": "피사체를 화면 중앙에 배치",
    "Golden Ratio (황금 비율)": "황금비에 맞춘 자연스러운 구도",
    "Leading Lines (시선 유도선)": "선을 이용해 시선 유도",
    "Frame within a Frame (프레임 속 프레임)": "창문 등으로 이중 프레임 구성",
    "Symmetry (대칭)": "좌우 또는 상하 대칭 구도",
    "Diagonal Composition (대각선 구도)": "대각선을 활용한 역동적 구도",
    
    # 카메라 무빙 (영상)
    "Pan Right (오른쪽 팬)": "카메라가 오른쪽으로 회전",
    "Pan Left (왼쪽 팬)": "카메라가 왼쪽으로 회전",
    "Tilt Up (위로 틸트)": "카메라가 위로 기울임",
    "Tilt Down (아래로 틸트)": "카메라가 아래로 기울임",
    "Zoom In (줌 인)": "피사체를 향해 줌 확대",
    "Zoom Out (줌 아웃)": "피사체에서 멀어지며 줌 축소",
    "Dolly In (돌리 인)": "카메라가 피사체 쪽으로 이동",
    "Dolly Out (돌리 아웃)": "카메라가 피사체에서 멀어짐",
    "Tracking Shot (트래킹 샷)": "피사체를 따라가며 촬영",
    "Crane Shot (크레인 샷)": "크레인으로 위아래 이동하며 촬영",
    "Steadicam (스테디캠)": "부드럽게 이동하며 흔들림 없는 촬영",
    "Handheld (핸드헬드)": "손으로 들고 찍어 현장감 있는 촬영",
    "Orbit Shot (카메라 회전)": "피사체 주위를 돌며 촬영",
    "Follow Shot (팔로우 샷)": "움직이는 피사체를 따라가는 샷",
    
    # 효과
    "Motion Blur (모션 블러)": "움직임에 의한 잔상 효과",
    "Slow Motion (슬로우 모션)": "느린 속도로 재생되는 효과",
    "Film Grain (필름 그레인)": "필름 카메라의 입자감 효과",
    "Bloom (블룸)": "밝은 부분이 번지는 효과",
    "Light Leaks (라이트 리크)": "빛이 새어 들어오는 효과",
    "Glitch (글리치)": "디지털 오류 같은 왜곡 효과",
    "Time-Lapse (타임랩스)": "시간을 압축해서 보여주는 효과",
    
    # 포즈
    "portrait (초상화)": "얼굴과 상체가 주가 되는 인물 사진",
    "Side profile (측면 초상화)": "옆모습을 담은 인물 사진",
    "upper body (상반신)": "허리 위 상체만 보이는 구도",
    "full body (전신)": "머리부터 발끝까지 전신 구도",
    
    # 표정
    "Smiling (미소)": "밝고 자연스러운 미소 표정",
    "Serious (진지한)": "진지하고 무표정한 얼굴",
    "Mysterious (신비로운)": "알 수 없는 신비로운 표정",
    "Confident (자신감 있는)": "당당하고 자신감 넘치는 표정",
    "Calm (차분한)": "평온하고 차분한 표정",
    
    # 품질
    "8K": "8K 해상도의 초고화질",
    "4K": "4K 해상도의 고화질",
    "Photorealistic (포토리얼리스틱)": "실제 사진처럼 사실적인 품질",
    "Hyperdetailed (하이퍼디테일드)": "극도로 세밀한 디테일",
    "Sharp Focus (샤프 포커스)": "선명하고 또렷한 초점",
    "HDR (High Dynamic Range)": "넓은 명암 범위의 이미지",
    
    # 렌더링
    "Unreal Engine (언리얼 엔진)": "언리얼 엔진 스타일의 3D 렌더링",
    "Octane Render (옥테인 렌더)": "옥테인 렌더러 스타일",
    "V-Ray": "V-Ray 렌더러의 사실적 렌더링",
    "cel shading": "셀 애니메이션 스타일의 렌더링",
}

def main():
    print("Adding option descriptions to Excel...")
    
    # Create descriptions dataframe
    data = [{"option_id": k, "description": v} for k, v in DESCRIPTIONS.items()]
    desc_df = pd.DataFrame(data)
    
    # Load existing workbook and add new sheet
    wb = load_workbook(EXCEL_PATH)
    
    # Remove existing sheet if exists
    if "옵션설명" in wb.sheetnames:
        del wb["옵션설명"]
    
    # Create new sheet
    ws = wb.create_sheet("옵션설명")
    
    # Add headers
    ws.cell(row=1, column=1, value="option_id")
    ws.cell(row=1, column=2, value="description")
    
    # Add data
    for idx, (k, v) in enumerate(DESCRIPTIONS.items(), start=2):
        ws.cell(row=idx, column=1, value=k)
        ws.cell(row=idx, column=2, value=v)
    
    # Save
    wb.save(EXCEL_PATH)
    print(f"Added {len(DESCRIPTIONS)} descriptions to '옵션설명' sheet")
    print("Now run: npm run convert:data")

if __name__ == "__main__":
    main()
