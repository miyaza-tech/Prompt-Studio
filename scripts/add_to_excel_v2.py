#!/usr/bin/env python3
"""
prompt_master_v2.xlsx의 Token_DB 시트에 프롬프트 해부학 옵션 12개를 추가.
Token_Input 시트에도 동기화 추가.
Group_DB에 Skin Realism 그룹 추가.
"""

import openpyxl
import re
import shutil
from pathlib import Path

BASE_DIR = Path(__file__).parent.parent
EXCEL_PATH = BASE_DIR / "data" / "prompt_master_v2.xlsx"

# Backup
backup_path = EXCEL_PATH.with_suffix('.xlsx.pre-anatomy-v2-backup')
shutil.copy2(EXCEL_PATH, backup_path)
print(f"Backup: {backup_path.name}")

wb = openpyxl.load_workbook(EXCEL_PATH)

# ============================================================
# Helper functions
# ============================================================
def make_token_key(label_short):
    """Generate token_key from label_short."""
    s = label_short.lower()
    s = re.sub(r'[^a-z0-9\s]', '', s)
    s = re.sub(r'\s+', '_', s.strip())
    return s[:50]

def make_token_id(category_key, group_key, token_key):
    return f"{category_key}__{group_key}__{token_key}"

# ============================================================
# Define new options
# ============================================================
NEW_OPTIONS = [
    # 조명 / 자연광
    {
        "category": "조명",
        "category_key": "cat_ecfc9e6b",
        "group": "자연광",
        "group_key": "default",
        "label_short": "Harsh Midday Sun",
        "description": "정오의 강한 직사광선으로 깊은 그림자와 높은 대비를 만듦. 다큐멘터리나 거친 사실적 묘사에 적합",
        "raw": "Harsh Midday Sun (정오의 강한 직사광선으로 깊은 그림자와 높은 대비를 만듦. 다큐멘터리나 거친 사실적 묘사에 적합)",
    },
    {
        "category": "조명",
        "category_key": "cat_ecfc9e6b",
        "group": "자연광",
        "group_key": "default",
        "label_short": "Late Afternoon Warmth",
        "description": "늦은 오후의 따뜻한 빛으로 긴 그림자와 부드러운 금빛 톤을 연출. 감성적인 인물 사진에 효과적",
        "raw": "Late Afternoon Warmth (늦은 오후의 따뜻한 빛으로 긴 그림자와 부드러운 금빛 톤을 연출. 감성적인 인물 사진에 효과적)",
    },
    {
        "category": "조명",
        "category_key": "cat_ecfc9e6b",
        "group": "자연광",
        "group_key": "default",
        "label_short": "Soft Diffused Daylight",
        "description": "부드럽게 확산된 자연광으로 그림자가 거의 없는 균일한 조명. 제품 촬영이나 자연스러운 인물 사진에 적합",
        "raw": "Soft Diffused Daylight (부드럽게 확산된 자연광으로 그림자가 거의 없는 균일한 조명. 제품 촬영이나 자연스러운 인물 사진에 적합)",
    },
    # 카메라&렌즈 / 스트리트
    {
        "category": "카메라&렌즈",
        "category_key": "cat_e7be0e5f",
        "group": "스트리트",
        "group_key": "default",
        "label_short": "Leica M11 + Summilux 50mm f/1.4",
        "description": "50mm 화각의 클래식한 렌즈. 보케가 아름답고 인물·스트리트 모두에서 뛰어난 표현력",
        "raw": "Leica M11 + Summilux 50mm f/1.4(50mm 화각의 클래식한 렌즈. 보케가 아름답고 인물·스트리트 모두에서 뛰어난 표현력)",
    },
    # 카메라&렌즈 / 매크로
    {
        "category": "카메라&렌즈",
        "category_key": "cat_e7be0e5f",
        "group": "매크로",
        "group_key": "default",
        "label_short": "Nikon Z9 + NIKKOR Z MC 105mm f/2.8 VR S",
        "description": "1:1 매크로에 VR 보정으로 손떨림 방지. 곤충, 꽃 등 극접사 촬영의 끝판왕",
        "raw": "Nikon Z9 + NIKKOR Z MC 105mm f/2.8 VR S(1:1 매크로에 VR 보정으로 손떨림 방지. 곤충, 꽃 등 극접사 촬영의 끝판왕)",
    },
    # 품질 & 디테일 / Skin Realism
    {
        "category": "품질 & 디테일",
        "category_key": "cat_f0183f05",
        "group": "Skin Realism",
        "group_key": "skin_realism",
        "label_short": "Visible Pores",
        "description": "피부의 모공이 자연스럽게 보이는 극사실적 표현",
        "raw": "Visible Pores (피부의 모공이 자연스럽게 보이는 극사실적 표현)",
    },
    {
        "category": "품질 & 디테일",
        "category_key": "cat_f0183f05",
        "group": "Skin Realism",
        "group_key": "skin_realism",
        "label_short": "Natural Skin Oils",
        "description": "피부 위의 자연스러운 유분기와 윤기 표현",
        "raw": "Natural Skin Oils (피부 위의 자연스러운 유분기와 윤기 표현)",
    },
    {
        "category": "품질 & 디테일",
        "category_key": "cat_f0183f05",
        "group": "Skin Realism",
        "group_key": "skin_realism",
        "label_short": "Fine Facial Hair",
        "description": "잔털과 솜털이 보이는 세밀한 피부 묘사",
        "raw": "Fine Facial Hair (잔털과 솜털이 보이는 세밀한 피부 묘사)",
    },
    {
        "category": "품질 & 디테일",
        "category_key": "cat_f0183f05",
        "group": "Skin Realism",
        "group_key": "skin_realism",
        "label_short": "Unretouched",
        "description": "보정하지 않은 원본 그대로의 피부 질감",
        "raw": "Unretouched (보정하지 않은 원본 그대로의 피부 질감)",
    },
    {
        "category": "품질 & 디테일",
        "category_key": "cat_f0183f05",
        "group": "Skin Realism",
        "group_key": "skin_realism",
        "label_short": "Raw Authentic Beauty",
        "description": "있는 그대로의 자연스러운 아름다움 표현",
        "raw": "Raw Authentic Beauty (있는 그대로의 자연스러운 아름다움 표현)",
    },
    {
        "category": "품질 & 디테일",
        "category_key": "cat_f0183f05",
        "group": "Skin Realism",
        "group_key": "skin_realism",
        "label_short": "Detailed Skin Texture",
        "description": "피부 결, 주름, 잡티 등 세밀한 질감 묘사",
        "raw": "Detailed Skin Texture (피부 결, 주름, 잡티 등 세밀한 질감 묘사)",
    },
    {
        "category": "품질 & 디테일",
        "category_key": "cat_f0183f05",
        "group": "Skin Realism",
        "group_key": "skin_realism",
        "label_short": "Raw Beauty Documentation",
        "description": "자연미를 기록하는 다큐멘터리적 접근",
        "raw": "Raw Beauty Documentation (자연미를 기록하는 다큐멘터리적 접근)",
    },
]

# ============================================================
# 1. Add to Token_DB
# ============================================================
ws = wb['Token_DB']
max_row = ws.max_row
added_count = 0

# Check for duplicates
existing_labels = set()
for r in range(2, max_row + 1):
    v = ws.cell(r, 7).value  # token_label_short
    if v:
        existing_labels.add(v)

for opt in NEW_OPTIONS:
    if opt["label_short"] in existing_labels:
        print(f"  SKIP (duplicate): {opt['label_short']}")
        continue
    
    token_key = make_token_key(opt["label_short"])
    token_id = make_token_id(opt["category_key"], opt["group_key"], token_key)
    
    new_row = max_row + 1 + added_count
    ws.cell(new_row, 1, token_id)           # token_id
    ws.cell(new_row, 2, "")                 # token_id_old
    ws.cell(new_row, 3, opt["category"])    # category
    ws.cell(new_row, 4, opt["category_key"])# category_key
    ws.cell(new_row, 5, opt["group"])       # group
    ws.cell(new_row, 6, opt["group_key"])   # group_key
    ws.cell(new_row, 7, opt["label_short"]) # token_label_short
    ws.cell(new_row, 8, opt["label_short"]) # token_label
    ws.cell(new_row, 9, opt["label_short"]) # display_label
    ws.cell(new_row, 10, opt["label_short"])# token_value
    ws.cell(new_row, 11, opt["description"])# description
    ws.cell(new_row, 12, opt["raw"])        # raw
    ws.cell(new_row, 13, "anatomy_script")  # source_sheet
    ws.cell(new_row, 14, new_row)           # source_row
    ws.cell(new_row, 15, 0)                 # source_col
    ws.cell(new_row, 16, token_key)         # token_key
    
    added_count += 1
    print(f"  Token_DB: {opt['label_short']}")

print(f"\nToken_DB: Added {added_count} rows (total: {max_row + added_count})")

# ============================================================
# 2. Add Skin Realism group to Group_DB
# ============================================================
ws_g = wb['Group_DB']
g_max = ws_g.max_row

# Check if Skin Realism already exists
skin_exists = False
for r in range(2, g_max + 1):
    if ws_g.cell(r, 2).value and 'Skin Realism' in str(ws_g.cell(r, 2).value):
        skin_exists = True
        break

if not skin_exists:
    new_r = g_max + 1
    ws_g.cell(new_r, 1, "품질 & 디테일")     # category
    ws_g.cell(new_r, 2, "Skin Realism")       # group
    ws_g.cell(new_r, 3, "피부 리얼리즘")      # group_display_name
    ws_g.cell(new_r, 4, "프롬프트 해부학 추가")# notes
    ws_g.cell(new_r, 5, 2)                    # sort_order
    print(f"\nGroup_DB: Added Skin Realism group")
else:
    print(f"\nGroup_DB: Skin Realism already exists")

# ============================================================
# 3. Add to Token_Input (for future reference)
# ============================================================
ws_i = wb['Token_Input']
i_max = ws_i.max_row
ti_added = 0

existing_input_names = set()
for r in range(2, i_max + 1):
    v = ws_i.cell(r, 5).value  # name col
    if v:
        existing_input_names.add(v)

for opt in NEW_OPTIONS:
    if opt["label_short"] in existing_input_names:
        continue
    new_r = i_max + 1 + ti_added
    ws_i.cell(new_r, 1, None)               # status
    ws_i.cell(new_r, 2, "프롬프트 해부학")   # notes
    ws_i.cell(new_r, 3, opt["category"])     # category
    ws_i.cell(new_r, 4, opt["group"])        # group
    ws_i.cell(new_r, 5, opt["label_short"])  # name
    ws_i.cell(new_r, 6, opt["description"])  # description
    ws_i.cell(new_r, 7, opt["label_short"])  # token_value
    ws_i.cell(new_r, 8, opt["label_short"])  # display_label
    ti_added += 1

print(f"Token_Input: Added {ti_added} rows")

# ============================================================
# Save
# ============================================================
wb.save(EXCEL_PATH)
print(f"\n=== Saved: {EXCEL_PATH.name} ===")
print("Now run: npm run convert:data")
